from __future__ import annotations

import pytest
import logging

from app.deepdoc.models import ParseOptions, SourceLocator
from app.deepdoc.parsers.base import ParseError
from app.deepdoc.parsers.pdf import PdfParser
from app.deepdoc.registry import ParserRegistry
from app.deepdoc.service import parse
from app.services.ingestion_service import _chunk_markdown_page


logger = logging.getLogger(__name__)


def test_parser_registry_resolves_by_mime_type_and_extension() -> None:
    registry = ParserRegistry()

    assert registry.resolve(mime_type="application/pdf", file_name="unknown").__class__.__name__ == "PdfParser"
    assert registry.resolve(mime_type=None, file_name="report.docx").__class__.__name__ == "DocxParser"
    assert registry.resolve(mime_type=None, file_name="sheet.xlsx").__class__.__name__ == "ExcelParser"
    assert registry.resolve(mime_type=None, file_name="deck.pptx").__class__.__name__ == "PptParser"
    assert registry.resolve(mime_type="text/markdown", file_name="doc.bin").__class__.__name__ == "MarkdownParser"
    assert registry.resolve(mime_type=None, file_name="notes.md").__class__.__name__ == "MarkdownParser"
    assert registry.resolve(mime_type="text/plain", file_name="doc.bin").__class__.__name__ == "TextParser"


def test_source_locator_serializes_without_empty_fields() -> None:
    locator = SourceLocator(type="excel", sheet="Sheet1", row_start=1, row_end=3, col_start=1, col_end=2)

    assert locator.to_dict() == {
        "type": "excel",
        "sheet": "Sheet1",
        "row_start": 1,
        "row_end": 3,
        "col_start": 1,
        "col_end": 2,
    }


@pytest.mark.asyncio
async def test_text_parser_returns_uniform_document(tmp_path) -> None:
    file_path = tmp_path / "note.txt"
    file_path.write_text("AI Second Brain\nmust cite sources", encoding="utf-8")

    document = await parse(
        file_path=str(file_path),
        mime_type="text/plain",
        file_name="note.txt",
        options=ParseOptions(enable_ocr=False),
    )

    assert document.text == "AI Second Brain\nmust cite sources"
    assert document.word_count == 6
    assert document.pages[0].page_number is None
    assert document.pages[0].locator is not None
    assert document.pages[0].locator.to_dict() == {"type": "text", "start": 0, "end": 33}
    assert document.metadata["ocr_used"] is False


@pytest.mark.asyncio
async def test_markdown_parser_returns_blocks_and_locator(tmp_path) -> None:
    file_path = tmp_path / "note.md"
    content = "# Deepdoc\n\nMarkdown paragraph\n## Details\n- cited item"
    file_path.write_text(content, encoding="utf-8")

    document = await parse(
        file_path=str(file_path),
        mime_type="text/markdown",
        file_name="note.md",
        options=ParseOptions(enable_ocr=False),
    )

    assert document.metadata["parser"] == "markdown"
    assert document.text == content
    assert document.pages[0].page_number is None
    assert document.pages[0].locator is not None
    assert document.pages[0].locator.to_dict() == {"type": "markdown", "start": 0, "end": len(content)}
    assert [block.kind for block in document.blocks] == ["title", "paragraph", "title", "paragraph"]
    assert document.blocks[0].text == "# Deepdoc"
    assert document.blocks[0].metadata == {"level": 1, "markdown_type": "heading", "path": ["Deepdoc"]}
    assert document.blocks[0].locator is not None
    assert document.blocks[0].locator.to_dict() == {"type": "markdown", "start": 0, "end": 9}
    assert document.blocks[3].text == "- cited item"
    assert document.blocks[3].metadata["markdown_type"] == "list"


@pytest.mark.asyncio
async def test_markdown_parser_preserves_structural_blocks_and_locators(tmp_path) -> None:
    file_path = tmp_path / "structured.md"
    content = """# Guide

Intro paragraph spans
two source lines.

- first item
- second item
  continued detail

```python
def hello():
    return "world"
```

| Name | Value |
| --- | --- |
| A | 1 |

<Warning title="注意">
Keep this warning together.
</Warning>
"""
    file_path.write_text(content, encoding="utf-8")

    document = await parse(
        file_path=str(file_path),
        mime_type="text/markdown",
        file_name="structured.md",
        options=ParseOptions(enable_ocr=False),
    )

    markdown_types = [block.metadata.get("markdown_type") for block in document.blocks]

    assert markdown_types == ["heading", "paragraph", "list", "code", "table", "html_block"]
    assert "- first item\n- second item\n  continued detail" in document.blocks[2].text
    assert 'return "world"' in document.blocks[3].text
    assert document.blocks[4].kind == "table"
    assert document.blocks[5].metadata["tag"] == "Warning"
    for block in document.blocks:
        assert block.locator is not None
        assert content[block.locator.start : block.locator.end] == block.text


@pytest.mark.asyncio
async def test_markdown_parser_recognizes_front_matter_setext_and_blockquote(tmp_path) -> None:
    file_path = tmp_path / "structured-extra.md"
    content = """---
title: Import Guide
---

Import Guide
============

> Keep this warning with its quote.
> It explains a single caution.

Body text.
"""
    file_path.write_text(content, encoding="utf-8")

    document = await parse(
        file_path=str(file_path),
        mime_type="text/markdown",
        file_name="structured-extra.md",
        options=ParseOptions(enable_ocr=False),
    )

    markdown_types = [block.metadata.get("markdown_type") for block in document.blocks]

    assert markdown_types == ["front_matter", "heading", "blockquote", "paragraph"]
    assert document.blocks[1].kind == "title"
    assert document.blocks[1].metadata == {"level": 1, "markdown_type": "heading", "path": ["Import Guide"]}
    assert document.blocks[2].metadata["path"] == ["Import Guide"]
    for block in document.blocks:
        assert block.locator is not None
        assert content[block.locator.start : block.locator.end] == block.text


@pytest.mark.asyncio
async def test_markdown_chunk_builder_keeps_blocks_and_heading_context(tmp_path) -> None:
    file_path = tmp_path / "chunks.md"
    content = """# Product

## Setup

Intro text.

- install SDK
- configure credentials

```cpp
engine->enableCustomVideoCapture(true, &captureConfig);
```

| Type | API |
| --- | --- |
| Raw | sendCustomVideoCaptureRawData |

## Troubleshooting

""" + "\n\n".join(f"Paragraph {index} " + ("detail " * 20) for index in range(8))
    file_path.write_text(content, encoding="utf-8")

    document = await parse(
        file_path=str(file_path),
        mime_type="text/markdown",
        file_name="chunks.md",
        options=ParseOptions(enable_ocr=False),
    )

    chunks = _chunk_markdown_page(document.pages[0], chunk_size=350)

    assert len(chunks) > 1
    assert chunks[0].content.startswith("# Product\n\n## Setup")
    assert "- install SDK\n- configure credentials" in chunks[0].content
    assert "engine->enableCustomVideoCapture" in chunks[0].content
    assert "| Raw | sendCustomVideoCaptureRawData |" in chunks[0].content
    assert all(
        chunk.locator is not None
        and chunk.locator["type"] == "markdown"
        and chunk.locator["start"] == chunk.start_pos
        and chunk.locator["end"] == chunk.end_pos
        for chunk in chunks
    )
    assert chunks[0].locator is not None
    assert chunks[0].locator["path"] == ["Product", "Setup"]
    assert any(chunk.content.startswith("# Product\n\n## Troubleshooting") for chunk in chunks[1:])


@pytest.mark.asyncio
async def test_markdown_chunk_builder_splits_long_paragraphs_without_losing_locator(tmp_path) -> None:
    file_path = tmp_path / "long.md"
    long_paragraph = " ".join(f"Sentence {index} has enough detail." for index in range(40))
    content = f"# Long Section\n\n{long_paragraph}"
    file_path.write_text(content, encoding="utf-8")

    document = await parse(
        file_path=str(file_path),
        mime_type="text/markdown",
        file_name="long.md",
        options=ParseOptions(enable_ocr=False),
    )

    chunks = _chunk_markdown_page(document.pages[0], chunk_size=260)

    assert len(chunks) > 2
    assert all(chunk.content.startswith("# Long Section") for chunk in chunks)
    assert all(len(chunk.content) <= 260 for chunk in chunks)
    for chunk in chunks:
        assert chunk.locator == {"type": "markdown", "start": chunk.start_pos, "end": chunk.end_pos, "path": ["Long Section"]}
        assert content[chunk.start_pos : chunk.end_pos].strip() in chunk.content


@pytest.mark.asyncio
async def test_markdown_chunk_builder_handles_real_custom_video_capture_fixture() -> None:
    file_path = "tests/test_files/custom-video-capture.md"
    # print(file_path)
    # logger.debug(f"file_path:{file_path}")
    document = await parse(
        file_path=file_path,
        mime_type="text/markdown",
        file_name="custom-video-capture.md",
        options=ParseOptions(enable_ocr=False),
    )
    print(f"Document page len:{len(document.pages)}")
    print(f"Document text:\n{document.pages[0].text}\n")
    chunks = _chunk_markdown_page(document.pages[0])
    # logger.info(f"Generated {len(chunks)} chunks from markdown with code, table, and HTML blocks.")
    # for i, chunk in enumerate(chunks):
    #     # logger.info(f"Chunk {i}: type={chunk.type}, content length={len(chunk.content)}")
    #     logger.debug(f"Chunk {i} content preview: {chunk.content[:200]}...")


    assert chunks
    assert all(
        chunk.locator is not None
        and chunk.locator["type"] == "markdown"
        and chunk.locator["start"] == chunk.start_pos
        and chunk.locator["end"] == chunk.end_pos
        for chunk in chunks
    )
    assert any("```mermaid" in chunk.content and "sequenceDiagram" in chunk.content for chunk in chunks)
    assert any("<Warning title=\"注意\">" in chunk.content and "</Warning>" in chunk.content for chunk in chunks)
    assert any("|视频帧类型|bufferType|发送视频帧数据接口|" in chunk.content for chunk in chunks)


@pytest.mark.asyncio
async def test_image_parser_requires_ocr_when_disabled(tmp_path) -> None:
    file_path = tmp_path / "scan.png"
    file_path.write_bytes(b"not a real image")

    with pytest.raises(ParseError, match="requires OCR"):
        await parse(
            file_path=str(file_path),
            mime_type="image/png",
            file_name="scan.png",
            options=ParseOptions(enable_ocr=False),
        )


@pytest.mark.asyncio
async def test_pdf_parser_returns_page_based_document(tmp_path) -> None:
    from pypdf import PdfWriter

    file_path = tmp_path / "sample.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with file_path.open("wb") as fp:
        writer.write(fp)

    document = await PdfParser().parse(
        file_path=str(file_path),
        mime_type="application/pdf",
        file_name="sample.pdf",
        options=ParseOptions(enable_ocr=False),
    )

    assert document.metadata["parser"] == "pdf"
    assert document.metadata["ocr_used"] is False
    assert len(document.pages) == 1
    assert document.pages[0].page_number == 1
    assert document.pages[0].locator is not None
    assert document.pages[0].locator.to_dict() == {"type": "pdf", "page": 1, "start": 0, "end": 0}


@pytest.mark.asyncio
async def test_docx_parser_returns_paragraph_locator(tmp_path) -> None:
    from docx import Document

    file_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Deepdoc supports DOCX")
    doc.add_paragraph("Each paragraph should be locatable")
    doc.save(file_path)

    document = await parse(
        file_path=str(file_path),
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_name="sample.docx",
        options=ParseOptions(enable_ocr=False),
    )

    assert document.metadata["parser"] == "docx"
    assert document.metadata["ocr_used"] is False
    assert "Deepdoc supports DOCX" in document.text
    assert len(document.pages) == 1
    assert document.pages[0].locator is not None
    assert document.pages[0].locator.to_dict() == {"type": "docx", "paragraph_start": 0, "paragraph_end": 1}
    assert len(document.blocks) == 2


@pytest.mark.asyncio
async def test_excel_parser_returns_sheet_locator(tmp_path) -> None:
    from openpyxl import Workbook

    file_path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A1"] = "Deepdoc"
    sheet["B1"] = "Excel"
    sheet["A2"] = "Parser"
    workbook.save(file_path)

    document = await parse(
        file_path=str(file_path),
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        file_name="sample.xlsx",
        options=ParseOptions(enable_ocr=False),
    )

    assert document.metadata["parser"] == "excel"
    assert document.metadata["ocr_used"] is False
    assert "Deepdoc" in document.text
    assert len(document.pages) == 1
    assert document.pages[0].locator is not None
    assert document.pages[0].locator.to_dict() == {
        "type": "excel",
        "sheet": "Sheet1",
        "row_start": 1,
        "row_end": 2,
        "col_start": 1,
        "col_end": 2,
    }


@pytest.mark.asyncio
async def test_ppt_parser_returns_slide_locator(tmp_path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    file_path = tmp_path / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text = "Deepdoc PPT template"
    presentation.save(file_path)

    document = await parse(
        file_path=str(file_path),
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        file_name="sample.pptx",
        options=ParseOptions(enable_ocr=False),
    )

    assert document.metadata["parser"] == "ppt"
    assert document.metadata["ocr_used"] is False
    assert "Deepdoc PPT template" in document.text
    assert len(document.pages) >= 1
    assert document.pages[0].locator is not None
    assert document.pages[0].locator.to_dict() == {"type": "ppt", "page": 1, "slide": 1, "start": 0, "end": 20}
