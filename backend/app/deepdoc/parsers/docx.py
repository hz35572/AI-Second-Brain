from __future__ import annotations

from app.deepdoc.models import ParseOptions, ParsedBlock, ParsedDocument, ParsedPage, SourceLocator
from app.deepdoc.parsers.base import BaseParser


class DocxParser(BaseParser):
    mime_types = {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    }
    extensions = {".docx", ".doc"}

    async def parse(
        self, *, file_path: str, mime_type: str | None, file_name: str, options: ParseOptions
    ) -> ParsedDocument:
        try:
            from docx import Document  # type: ignore
        except Exception as exc:  # noqa: BLE001
            raise RuntimeError("DOCX parser dependency missing: install python-docx") from exc

        document = Document(file_path)
        blocks: list[ParsedBlock] = []
        parts: list[str] = []
        for index, paragraph in enumerate(document.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            locator = SourceLocator(type="docx", paragraph_start=index, paragraph_end=index)
            blocks.append(ParsedBlock(text=text, kind="paragraph", locator=locator))
            parts.append(text)

        full_text = "\n".join(parts)
        page_locator = SourceLocator(type="docx", paragraph_start=0, paragraph_end=max(0, len(document.paragraphs) - 1))
        page = ParsedPage(text=full_text, page_number=None, locator=page_locator, blocks=blocks)
        return ParsedDocument(
            pages=[page],
            blocks=blocks,
            metadata={"parser": "docx", "mime_type": mime_type, "file_name": file_name, "ocr_used": False},
        )
