from __future__ import annotations

from app.deepdoc.models import ParseOptions, ParsedBlock, ParsedDocument, ParsedPage, SourceLocator
from app.deepdoc.parsers.base import BaseParser


class ExcelParser(BaseParser):
    mime_types = {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    }
    extensions = {".xlsx", ".xls"}

    async def parse(
        self, *, file_path: str, mime_type: str | None, file_name: str, options: ParseOptions
    ) -> ParsedDocument:
        try:
            from openpyxl import load_workbook  # type: ignore
        except Exception as exc:  # noqa: BLE001
            raise RuntimeError("Excel parser dependency missing: install openpyxl") from exc

        workbook = load_workbook(file_path, read_only=True, data_only=True)
        pages: list[ParsedPage] = []
        blocks: list[ParsedBlock] = []
        for sheet in workbook.worksheets:
            rows: list[str] = []
            max_row = sheet.max_row or 0
            max_col = sheet.max_column or 0
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(value.strip() for value in values):
                    rows.append("\t".join(values).rstrip())
            text = "\n".join(rows)
            locator = SourceLocator(
                type="excel",
                sheet=sheet.title,
                row_start=1 if max_row else None,
                row_end=max_row or None,
                col_start=1 if max_col else None,
                col_end=max_col or None,
            )
            block = ParsedBlock(text=text, kind="sheet", locator=locator)
            blocks.append(block)
            pages.append(ParsedPage(text=text, page_number=None, locator=locator, blocks=[block]))

        workbook.close()
        return ParsedDocument(
            pages=pages,
            blocks=blocks,
            metadata={"parser": "excel", "mime_type": mime_type, "file_name": file_name, "ocr_used": False},
        )
