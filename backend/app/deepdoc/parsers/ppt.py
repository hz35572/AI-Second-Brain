from __future__ import annotations

from app.deepdoc.models import ParseOptions, ParsedDocument, ParsedPage, SourceLocator
from app.deepdoc.parsers.base import BaseParser


class PptParser(BaseParser):
    mime_types = {
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/vnd.ms-powerpoint",
    }
    extensions = {".pptx", ".ppt"}

    async def parse(
        self, *, file_path: str, mime_type: str | None, file_name: str, options: ParseOptions
    ) -> ParsedDocument:
        try:
            from pptx import Presentation  # type: ignore
        except Exception as exc:  # noqa: BLE001
            raise RuntimeError("PPT parser dependency missing: install python-pptx") from exc

        presentation = Presentation(file_path)
        pages: list[ParsedPage] = []
        for index, slide in enumerate(presentation.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    parts.append(shape.text)
            text = "\n".join(parts)
            locator = SourceLocator(type="ppt", page=index, slide=index, start=0, end=len(text))
            pages.append(ParsedPage(text=text, page_number=index, locator=locator))

        return ParsedDocument(
            pages=pages,
            metadata={"parser": "ppt", "mime_type": mime_type, "file_name": file_name, "ocr_used": False},
        )
